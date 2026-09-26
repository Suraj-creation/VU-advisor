"""Excel adapter. Reads raw cell values + provenance; no domain interpretation happens here.

Emits table_row SourceUnits with content["table"] in:
  offering        one course cell-block of a semester-spread grid (code, title, prereq, L, T, P, C)
  basket_spread   a basket label row of a semester-spread grid (with its "Fixed Min" credits)
  structure       one course line of a Struct_* sheet
  basket_struct   one basket line of a Struct_* sheet
  minor           one row of a minor table
  anomaly         anything the grid layout could not place (stray cell, odd formula) -> validation
"""
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl

from . import unit

SEM_SHEETS = {"Sem Spread BTECH-2022": 2022, "Sem_Spread_2023": 2023, "Sem_Spread_2024": 2024,
              "Sem_Spread_2025": 2025, "Sem_Spread_DS_2026": 2026}
STRUCT_SHEETS = {"Struct_2022": 2022, "Struct_2023": 2023, "Struct_2024": 2024,
                 "Struct_2025": 2025, "Struct_2026_DS": 2026}
# Semester blocks sit at fixed columns in every semester-spread sheet: code, title, prereq, L, T, P, C.
# (Header labels are unreliable -- some blocks lack the "Course Code" header.)
BLOCKS = {1: "E", 2: "L", 3: "T", 4: "AA", 5: "AI", 6: "AP", 7: "AX", 8: "BE"}
SUMMER_COLS = {"S", "AH", "AW"}
MINOR_NAMES = {"Law Minor": "Law", "Design Minor": "Design", "Psychology": "Psychology",
               "Economics": "Economics", "Finance": "Finance", "Marketing": "Marketing", "Start-up": "Start-up"}


def text(v):
    """Whitespace-normalised string, or None for empty / whitespace-only cells."""
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()
    return s or None


def number(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    return text(v)


def extract(path):
    wb = load_workbook(path)  # formulas kept (data_only=False) so anomalies are visible
    names = set(wb.sheetnames)
    if names & set(SEM_SHEETS):
        return semester_spread(wb, path) + structure(wb, path)
    if names & set(MINOR_NAMES):
        return minors(wb, path)
    raise ValueError(f"No known layout in {path}: {wb.sheetnames}")


def semester_spread(wb, path):
    out = []
    for sheet, batch in SEM_SHEETS.items():
        ws = wb[sheet]
        used, label, total_row = set(), None, None
        for r in range(3, ws.max_row + 1):
            a, b = text(ws.cell(r, 1).value), text(ws.cell(r, 2).value)
            if a and a.lower().startswith("total"):
                total_row = r
                break
            if b:  # basket label (top cell of a merged range); col A id is missing for B3 in 2026
                label = b
                out.append(unit("table_row", path, f"A{r}:D{r}", {
                    "table": "basket_spread", "batch": batch, "row": r, "basket_id": a, "label": b,
                    "fixed_min": number(ws.cell(r, 4).value)}, sheet=sheet))
            if label is None:
                continue
            for sem, col in BLOCKS.items():
                c0 = ci(col)
                raw = [ws.cell(r, c0 + i).value for i in range(7)]
                used.update((r, c0 + i) for i in range(7))
                if all(text(v) is None for v in raw):
                    continue
                out.append(unit("table_row", path, f"{col}{r}:{gl(c0 + 6)}{r}", {
                    "table": "offering", "batch": batch, "semester": sem, "row": r, "basket_label": label,
                    "code": raw[0], "title": raw[1], "prereq": raw[2],
                    "L": number(raw[3]), "T": number(raw[4]), "P": number(raw[5]), "C": number(raw[6])}, sheet=sheet))
        out += _grid_anomalies(ws, sheet, path, used, total_row)
    return out


def _grid_anomalies(ws, sheet, path, used, total_row):
    """Every non-empty cell must be a course block, a basket column, a total/label, or a sane row formula."""
    out = []
    for row in ws.iter_rows(min_row=3, max_row=total_row - 1 if total_row else ws.max_row):
        for c in row:
            v = c.value
            if text(v) is None or c.column <= 4 or (c.row, c.column) in used:
                continue
            col = gl(c.column)
            if col in SUMMER_COLS:
                out.append(_anomaly(path, sheet, c.coordinate, "value in SUMMER column", v))
            elif col in ("BL", "BM"):
                if isinstance(v, str) and v.startswith("="):
                    rows = {int(x) for x in re.findall(r"[A-Z]{1,2}(\d+)", v)}
                    if col == "BL" and rows != {c.row}:
                        out.append(_anomaly(path, sheet, c.coordinate, "row-total formula references another row", v))
                else:
                    out.append(_anomaly(path, sheet, c.coordinate, "row total is a hard-coded number, not a formula", v))
            else:
                out.append(_anomaly(path, sheet, c.coordinate, "stray cell outside the course grid", v))
    return out


def _anomaly(path, sheet, cell, what, value):
    return unit("table_row", path, cell, {"table": "anomaly", "what": what, "value": str(value)}, sheet=sheet)


def structure(wb, path):
    out = []
    for sheet, batch in STRUCT_SHEETS.items():
        ws = wb[sheet]
        for row in ws.iter_rows():
            for c in row:
                v = text(c.value)
                if not v:
                    continue
                if v == '"Baskets"':
                    out += _struct_baskets(ws, sheet, batch, path, c)
                right = text(ws.cell(c.row, c.column + 1).value)
                if right == "Cr" and ("Courses" in v or v in ("Electives", "Minor Total Credits")):
                    out += _struct_section(ws, sheet, batch, path, c, v)
    return out


def _struct_baskets(ws, sheet, batch, path, head):
    out, col = [], head.column
    for r in range(head.row + 1, ws.max_row + 1):
        label, cr = text(ws.cell(r, col).value), ws.cell(r, col + 1).value
        if not label:
            continue
        kind = "total" if label.lower().startswith("total") else "basket"
        if kind == "basket" and not isinstance(cr, (int, float)):
            continue  # "(Components)" sub-header
        out.append(unit("table_row", path, f"{gl(col)}{r}:{gl(col + 1)}{r}", {
            "table": "basket_struct", "batch": batch, "kind": kind, "label": label, "credits": cr}, sheet=sheet))
        if kind == "total":
            break
    return out


def _struct_section(ws, sheet, batch, path, head, section):
    out = []
    for r in range(head.row + 1, ws.max_row + 1):
        num, title, cr = text(ws.cell(r, head.column - 1).value), text(ws.cell(r, head.column).value), ws.cell(r, head.column + 1).value
        if num == "#" or "Total" in ((num or "") + (title or "")):
            break
        if title or cr is not None:
            out.append(unit("table_row", path, f"{gl(head.column - 1)}{r}:{gl(head.column + 1)}{r}", {
                "table": "structure", "batch": batch, "section": section, "number": num,
                "title": title, "credits": number(cr)}, sheet=sheet))
    return out


def minors(wb, path):
    out = []
    for sheet, minor in MINOR_NAMES.items():
        ws = wb[sheet]
        hdr = {text(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if text(ws.cell(1, c).value)}
        cc, tc, lc, crc, sc, pc = (hdr["Course Code"], hdr["Course Title"], hdr["L"], hdr["Credit"],
                                   hdr["Semester"], hdr["Pre-rq"])
        a1 = text(ws.cell(1, 1).value)  # some sheets put the first batch label on the header row
        batch = int(a1[:4]) if a1 and "Batch" in a1 else None
        for r in range(2, ws.max_row + 1):
            a = text(ws.cell(r, 1).value)
            if a and "Batch" in a:
                batch = int(a[:4])
            code, title = text(ws.cell(r, cc).value), text(ws.cell(r, tc).value)
            if not (code or title) or (title or "").endswith(" Minor") and not code:
                continue
            first = min(cc, tc)
            out.append(unit("table_row", path, f"{gl(first)}{r}:{gl(pc + 1)}{r}", {
                "table": "minor", "minor": minor, "batch": batch, "row": r,
                "code": ws.cell(r, cc).value, "title": ws.cell(r, tc).value,
                "L": number(ws.cell(r, lc).value), "T": number(ws.cell(r, lc + 1).value),
                "P": number(ws.cell(r, lc + 2).value), "C": number(ws.cell(r, crc).value),
                "semester": number(ws.cell(r, sc).value), "prereq": ws.cell(r, pc).value,
                "prereq_title": ws.cell(r, pc + 1).value if pc + 1 <= ws.max_column else None}, sheet=sheet))
    return out
